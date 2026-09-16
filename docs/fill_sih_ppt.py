"""Fill the SIH 2026 idea template with JanNidhi content (fill-in mode)."""
from pptx import Presentation
from pptx.util import Inches, Pt

SRC = r'C:\Users\poula\OneDrive\Desktop\SIH2026-IDEA-Presentation-Format.pptx'
OUT = r'D:\SIH 1\docs\JanNidhi_SIH2026_Idea_Presentation.pptx'

prs = Presentation(SRC)
slides = prs.slides


def set_body(slide, blocks, new_top=None, new_height=None):
    """Rebuild the first 'TextBox N' body: blocks = [(text, is_pointer)]."""
    box = [sh for sh in slide.shapes
           if sh.has_text_frame and sh.name.startswith('TextBox')][0]
    fname, fsize = None, None
    for p in box.text_frame.paragraphs:
        for r in p.runs:
            fname = r.font.name or fname
            fsize = r.font.size or fsize
            break
        if fname or fsize:
            break
    base = fsize.pt if fsize else 14
    content_sz = max(10, min(12, base - 2))
    if new_top is not None:
        box.top = Inches(new_top)
    if new_height is not None:
        box.height = Inches(new_height)
    if box.width.inches < 11:
        box.width = Inches(12.4)
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    first = True
    for text, is_pointer in blocks:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r = p.add_run()
        r.text = text
        r.font.name = fname or 'Calibri'
        if is_pointer:
            r.font.size = Pt(min(base, 14))
            r.font.bold = True
        else:
            r.font.size = Pt(content_sz)
        p.space_after = Pt(4)


# ---------- SLIDE 1 — TITLE PAGE ----------
s1 = slides[0]
tb = [sh for sh in s1.shapes if sh.has_text_frame and sh.name == 'TextBox 9'][0]
fields = [
    ('Problem Statement ID –', 'SIH26102'),
    ('Problem Statement Title-', 'MPLADS Audit & Anomaly Prioritization Platform'),
    ('Theme-', '(as registered on SIH portal)'),
    ('PS Category-', 'Software'),
    ('Team ID-', '(fill from portal)'),
    ('Team Name (Registered on portal)', '(fill from portal)'),
]
tf = tb.text_frame
first = True
for label, value in fields:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    r1 = p.add_run(); r1.text = label + ' '
    r2 = p.add_run(); r2.text = value
    for r in (r1, r2):
        r.font.name = 'Calibri'
        r.font.size = Pt(16)
    r1.font.bold = True
    p.space_after = Pt(10)

# ---------- SLIDE 2 — PROPOSED SOLUTION ----------
s2 = slides[1]
for sh in s2.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() == 'IDEA TITLE':
        sh.text_frame.paragraphs[0].runs[0].text = \
            'JANNIDHI – AI-DRIVEN MPLADS AUDIT & RISK PLATFORM'
set_body(s2, [
    ('Proposed Solution (Describe your Idea/Solution/Prototype)', True),
    ('– JanNidhi: AI audit-prioritization platform that risk-scores 100% of MPLADS works — 132,379 live works ingested from the official portal', False),
    ('– Unified score per work: 0.45 × 21-rule engine + 0.30 × Isolation Forest + 0.25 × XGBoost, ranked by likelihood × financial impact', False),
    ('– Evidence-grounded Case Packets: every flag explained, recommended action, one-click reviewer feedback loop', False),
    ('Detailed explanation of the proposed solution', True),
    ('– Live ingestion from the MPLADS portal REST API (nightly, 03:00 IST) + MP allocation ledger (774 MPs)', False),
    ('– 6 detection families: financial reconciliation · timeline · cost outliers (peer MAD) · vendor capture · ghost works (cross-MP NLP duplicates) · MP ceiling breaches', False),
    ('– Reviewer UI: priority queue, MP/state dossiers, MP comparison, open-data CSV export', False),
    ('How it addresses the problem', True),
    ('– Manual audits sample a fraction; JanNidhi ranks ALL works — 16,249 high-risk works flagged of 132k, top-ranked reviewed first', False),
    ('– Live watch-metrics: ₹11,698 Cr allocated · ₹3,998 Cr spent · ₹1,575 Cr ongoing-work payments · 88,000 pending works', False),
    ('Innovation and uniqueness', True),
    ('– First system to replicate the MPLADS dashboard\u2019s internal REST API for automated live sync (no official bulk export exists)', False),
    ('– Fully explainable flags (auditor sees WHY) + weak-supervision model retraining from reviewer outcomes', False),
], new_top=1.35, new_height=5.5)

# ---------- SLIDE 3 — TECHNICAL APPROACH ----------
set_body(slides[2], [
    ('Technologies to be used (e.g. programming languages, frameworks, hardware)', True),
    ('– React + Vite + Tailwind · Python FastAPI + SQLAlchemy · scikit-learn Isolation Forest · XGBoost · pandas + TF-IDF NLP · APScheduler · SQLite WAL (PostgreSQL-ready)', False),
    ('Methodology and process for implementation (Flow Charts/Images/ working prototype)', True),
    ('– MPLADS Portal REST API  >  Nightly Ingestion (132k works + 774 MP ceilings)  >  Reshape to work-level facts', False),
    ('– 21-Rule Engine (6 families: financial / timeline / cost / vendor / ghost-work / MP-ceiling)  >  ML scoring (Isolation Forest + XGBoost)', False),
    ('– Unified risk score  >  priority rank  >  Reviewer UI (priority queue / case packets / MP & state dossiers / compare)', False),
    ('– Working prototype live on real data: 146,446 rule triggers · 7,953 statistical anomalies · 16,249 high-risk works', False),
], new_top=1.6, new_height=5.2)

# ---------- SLIDE 4 — FEASIBILITY AND VIABILITY ----------
set_body(slides[3], [
    ('Analysis of the feasibility of the idea', True),
    ('– Fully working prototype validated against the live official portal (figures match the public dashboard exactly)', False),
    ('– Commodity stack, single small server, zero paid dependencies', False),
    ('Potential challenges and risks', True),
    ('– No official bulk API; portal response shape may change', False),
    ('– Very large payloads (90 MB+) from a slow government server', False),
    ('– No ground-truth fraud labels for supervised training', False),
    ('Strategies for overcoming these challenges', True),
    ('– Direct REST replication + raw-feed caching + offline re-scoring (engine improvements without re-fetching)', False),
    ('– Background nightly sync (03:00 IST) with WAL + chunked commits — site stays live during syncs', False),
    ('– Weak supervision from the rule engine + human reviewer feedback loop for model retraining', False),
], new_top=1.6, new_height=5.2)

# ---------- SLIDE 5 — IMPACT AND BENEFITS ----------
set_body(slides[4], [
    ('Potential impact on the target audience', True),
    ('– MoSPI auditors: audit triage from months to minutes — 132k works ranked by risk with evidence packets', False),
    ('– Citizens / journalists / researchers: public transparency dashboard + open-data CSV export', False),
    ('Benefits of the solution (social, economic, environmental, etc.)', True),
    ('– Social: protects public welfare funds; fully explainable flags — no black-box accusations', False),
    ('– Economic: early detection of vendor capture & ghost works safeguards ₹11,698 Cr allocated; ₹1,575 Cr ongoing-work payments under watch', False),
    ('– Governance: auditable sync_logs + review_logs chain; 44,379 completed vs 88,000 pending works tracked', False),
], new_top=1.6, new_height=5.2)

# ---------- SLIDE 6 — RESEARCH AND REFERENCES ----------
set_body(slides[5], [
    ('Details / Links of the reference and research work', True),
    ('– MPLADS Guidelines & portal — mplads.mospi.gov.in (data source; dashboard REST endpoint getTilesReportData)', False),
    ('– CAG of India audit reports on MPLADS (ghost works, vendor collusion, fund-diversion patterns)', False),
    ('– Empowered-Indian civic-tech project — open MPLADS data methodology (github.com/Empowered-Indian/empowered-indian)', False),
    ('– Liu et al., 2008 — Isolation Forest (anomaly detection); Chen & Guestrin, 2016 — XGBoost', False),
    ('– Median Absolute Deviation (Hampel) — robust cost-outlier estimation', False),
], new_top=1.6, new_height=5.2)

# ---------- delete slide 7 (instructions) ----------
sld = list(prs.slides._sldIdLst)[6]
prs.part.drop_rel(sld.get(
    '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'))
prs.slides._sldIdLst.remove(sld)

prs.save(OUT)
print('saved:', OUT, '| slides:', len(list(prs.slides)))
